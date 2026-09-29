"""Nettoie les montants des factures et les écrit dans un classeur Excel consolidé."""

import datetime as dt
import re
from pathlib import Path
from typing import NamedTuple

import openpyxl

COLUMNS = ["Document", "Date", "Client", "Type", "Net Worth", "VAT", "Gross Worth"]
NUMBER_FORMAT = "#,##0.00"
DATE_FORMAT = "DD/MM/YYYY"

# Code devise ou symbole en préfixe d'un montant (ex. "INR 1,676,976.00")
CURRENCY_PREFIX = re.compile(r"^(?:[A-Z]{3}\s*|[$€£]\s*)(?=[\d.,-])")


class InvoiceRow(NamedTuple):
    document: str
    date: dt.date
    client: str
    totals: list[str]  # cellules de la ligne de total : Type, Net Worth, VAT, Gross Worth


def clean_amount(text: str) -> float | str:
    """Retire la devise et les séparateurs de milliers ; garde le texte si ce n'est pas un nombre."""
    value = CURRENCY_PREFIX.sub("", text.replace("\n", " ").strip())
    try:
        return float(value.replace(",", ""))
    except ValueError:
        return value


def export_to_xlsx(rows: list[InvoiceRow], output_path: Path) -> None:
    """Écrit l'en-tête puis une ligne par facture."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Facture"
    ws.append(COLUMNS)

    for row in rows:
        ws.append([row.document, row.date, row.client, *map(clean_amount, row.totals)])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, float):
                cell.number_format = NUMBER_FORMAT
            elif isinstance(cell.value, dt.date):
                cell.number_format = DATE_FORMAT

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
